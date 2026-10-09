from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from app.core.dependencies import get_current_user, get_current_admin, get_current_teacher
from app.services.academic_service import result_service, attendance_service
from app.services.user_service import student_service
from app.core.database import get_collection
from bson import ObjectId
import io
from datetime import datetime

router = APIRouter(prefix="/reports", tags=["Reports"])


@router.get("/student/{student_id}/pdf")
async def generate_student_report_pdf(student_id: str, _=Depends(get_current_user)):
    """Generate student academic report as PDF"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch

    student = await student_service._get_student_with_details(student_id)
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    results = await result_service.get_student_results(student_id)
    attendance_summary = await attendance_service.get_attendance_summary(student_id)
    cgpa_data = await result_service.calculate_cgpa(student_id)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    title = Paragraph(f"<b>ACADEMIC REPORT CARD</b>", styles["Title"])
    elements.append(title)
    elements.append(Spacer(1, 0.2 * inch))

    # Student Info
    info_data = [
        ["Student Name", student.get("full_name", "N/A"), "CGPA", str(cgpa_data.get("cgpa", "N/A"))],
        ["Enrollment No.", student.get("enrollment_number", "N/A"), "Semester", str(student.get("semester", "N/A"))],
        ["Department", student.get("department_name", "N/A"), "Batch Year", str(student.get("batch_year", "N/A"))],
        ["Email", student.get("email", "N/A"), "Generated", datetime.now().strftime("%d-%m-%Y")],
    ]

    info_table = Table(info_data, colWidths=[1.5 * inch, 2.5 * inch, 1.2 * inch, 2 * inch])
    info_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.lightblue),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 0.3 * inch))

    # Results Table
    if results:
        elements.append(Paragraph("<b>Academic Results</b>", styles["Heading2"]))
        result_data = [["Subject", "Internal", "External", "Total", "Max", "Percentage", "Grade", "Status"]]
        for r in results:
            result_data.append([
                r.get("course_name", "N/A"),
                str(r.get("internal_marks", 0)),
                str(r.get("external_marks", 0)),
                str(r.get("total_marks", 0)),
                str(r.get("max_marks", 100)),
                f"{r.get('percentage', 0):.1f}%",
                r.get("grade", "N/A"),
                "PASS" if r.get("is_pass") else "FAIL",
            ])

        result_table = Table(result_data, colWidths=[2 * inch, 0.8 * inch, 0.8 * inch, 0.7 * inch, 0.7 * inch, 0.9 * inch, 0.6 * inch, 0.7 * inch])
        result_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("PADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.append(result_table)
        elements.append(Spacer(1, 0.2 * inch))

    # Attendance Summary
    if attendance_summary:
        elements.append(Paragraph("<b>Attendance Summary</b>", styles["Heading2"]))
        att_data = [["Subject", "Total Classes", "Present", "Absent", "Percentage", "Status"]]
        for a in attendance_summary:
            status = "Good" if a["percentage"] >= 75 else "Low"
            att_data.append([
                a.get("course_name", "N/A"),
                str(a.get("total", 0)),
                str(a.get("present", 0)),
                str(a.get("absent", 0)),
                f"{a.get('percentage', 0):.1f}%",
                status,
            ])

        att_table = Table(att_data, colWidths=[2.5 * inch, 1.2 * inch, 1 * inch, 1 * inch, 1 * inch, 0.8 * inch])
        att_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.darkgreen),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("PADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.append(att_table)

    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=report_{student_id}.pdf"},
    )


@router.get("/attendance/excel")
async def export_attendance_excel(
    course_id: str,
    date: str | None = None,
    from_date: str | None = None,
    to_date: str | None = None,
    current_user: dict = Depends(get_current_teacher),
):
    """Export attendance as Excel file"""
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment

    await attendance_service.get_manageable_course(course_id, current_user)
    records = await attendance_service.get_course_attendance(course_id, date, from_date, to_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    header_fill = PatternFill(start_color="1e3a5f", end_color="1e3a5f", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    headers = ["Enrollment Number", "Student Name", "Date", "Status", "Course"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row, record in enumerate(records, 2):
        ws.cell(row=row, column=1, value=record.get("enrollment_number", ""))
        ws.cell(row=row, column=2, value=record.get("student_name", ""))
        ws.cell(row=row, column=3, value=str(record.get("date", ""))[:10])
        ws.cell(row=row, column=4, value=record.get("status", ""))
        ws.cell(row=row, column=5, value=record.get("course_name", ""))

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 4

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=attendance_{course_id}_{date or 'history'}.xlsx"},
    )
