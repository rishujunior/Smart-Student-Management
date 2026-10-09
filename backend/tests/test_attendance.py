from datetime import date, datetime, timezone
from uuid import uuid4

import pytest
from bson import ObjectId
from httpx import AsyncClient

from app.core.database import get_collection
from app.core.security import create_access_token


@pytest.mark.asyncio
async def test_teacher_can_save_reload_and_update_daily_attendance(client: AsyncClient):
    departments = get_collection("departments")
    users = get_collection("users")
    teachers = get_collection("teachers")
    students = get_collection("students")
    courses = get_collection("courses")
    attendance = get_collection("attendance")

    department = await departments.find_one({"code": "CSE"})
    assert department, "CSE department must exist for the attendance integration test"

    suffix = uuid4().hex[:10]
    teacher_user_id = ObjectId()
    teacher_id = ObjectId()
    course_id = ObjectId()
    test_student_user_ids = [ObjectId(), ObjectId(), ObjectId()]
    test_student_ids = [ObjectId(), ObjectId(), ObjectId()]
    await users.insert_one({
        "_id": teacher_user_id,
        "email": f"attendance-teacher-{suffix}@test.sms.edu",
        "full_name": f"Attendance Teacher {suffix}",
        "role": "teacher",
        "is_active": True,
    })
    await teachers.insert_one({
        "_id": teacher_id,
        "user_id": str(teacher_user_id),
        "employee_id": f"AT{suffix}",
        "department_id": str(department["_id"]),
        "is_active": True,
    })
    await courses.insert_one({
        "_id": course_id,
        "name": f"Attendance Integration Course {suffix}",
        "code": f"AT{suffix}".upper(),
        "department_id": str(department["_id"]),
        "semester": 3,
        "teacher_id": str(teacher_id),
        "is_active": True,
        "subject_type": "theory",
        "created_at": datetime.now(timezone.utc),
    })

    for index, (user_id, student_id) in enumerate(zip(test_student_user_ids, test_student_ids, strict=True), 1):
        await users.insert_one({
            "_id": user_id,
            "email": f"attendance-student-{index}-{suffix}@test.sms.edu",
            "full_name": f"Attendance Test Student {index} {suffix}",
            "role": "student",
            "is_active": True,
        })
        await students.insert_one({
            "_id": student_id,
            "user_id": str(user_id),
            "enrollment_number": f"AT{suffix}{index}",
            "department_id": str(department["_id"]),
            "semester": 3,
            "is_active": True,
        })

    token = create_access_token({"sub": str(teacher_user_id), "role": "teacher"})
    headers = {"Authorization": f"Bearer {token}"}
    selected_date = date.today().isoformat()
    course_id_str = str(course_id)

    try:
        courses_response = await client.get("/api/v1/attendance/courses", headers=headers)
        assert courses_response.status_code == 200, courses_response.text
        assert any(course["id"] == course_id_str for course in courses_response.json())
        official_names = {
            "BCS301": "Data Structure",
            "BCS302": "Computer Organization & Architecture",
            "BCS303": "Discrete Structures & Theory of Logic",
            "BAS303": "Mathematics-IV",
            "BVE301": "Universal Human Values & Professional Ethics",
            "BCC302": "Python Programming",
            "BCS351": "Data Structure Lab",
            "BCS352": "Computer Organization & Architecture Lab",
            "BCS353": "Web Designing Workshop",
            "BCC351": "Internship Assessment / Mini Project",
        }
        official_courses = [course for course in courses_response.json() if course["code"] in official_names]
        assert [course["code"] for course in official_courses] == list(official_names)
        assert {course["code"]: course["name"] for course in official_courses} == official_names
        assert all("teacher_name" not in course and "faculty_name" not in course for course in courses_response.json())

        roster_response = await client.get(
            f"/api/v1/attendance/course/{course_id_str}/roster",
            params={"date": selected_date},
            headers=headers,
        )
        assert roster_response.status_code == 200, roster_response.text
        roster = roster_response.json()
        assert all(student_id in {student["id"] for student in roster["students"]} for student_id in map(str, test_student_ids))
        assert roster["summary"]["total_students"] == len(roster["students"])

        selected_students = [student for student in roster["students"] if student["id"] in set(map(str, test_student_ids))]
        records = [
            {"student_id": student["id"], "status": status}
            for student, status in zip(selected_students, ("present", "absent", "late"), strict=True)
        ]
        save_response = await client.post(
            "/api/v1/attendance",
            json={"course_id": course_id_str, "date": f"{selected_date}T00:00:00Z", "records": records},
            headers=headers,
        )
        assert save_response.status_code == 200, save_response.text
        assert save_response.json()["summary"]["present"] == 1
        assert save_response.json()["summary"]["absent"] == 1
        assert save_response.json()["summary"]["late"] == 1
        assert save_response.json()["summary"]["coverage"] == round(3 / len(roster["students"]) * 100, 2)

        reload_response = await client.get(
            f"/api/v1/attendance/course/{course_id_str}/roster",
            params={"date": selected_date},
            headers=headers,
        )
        reloaded = {student["id"]: student["attendance_status"] for student in reload_response.json()["students"]}
        assert all(reloaded[record["student_id"]] == record["status"] for record in records)
        historical_rates = {student["id"]: student["attendance_percentage"] for student in reload_response.json()["students"]}
        assert historical_rates[records[0]["student_id"]] == 100
        assert historical_rates[records[1]["student_id"]] == 0
        assert historical_rates[records[2]["student_id"]] == 100

        records[0]["status"] = "absent"
        update_response = await client.post(
            "/api/v1/attendance",
            json={"course_id": course_id_str, "date": f"{selected_date}T18:45:00Z", "records": records},
            headers=headers,
        )
        assert update_response.status_code == 200, update_response.text

        final_response = await client.get(
            f"/api/v1/attendance/course/{course_id_str}/roster",
            params={"date": selected_date},
            headers=headers,
        )
        final_statuses = {student["id"]: student["attendance_status"] for student in final_response.json()["students"]}
        final_rates = {student["id"]: student["attendance_percentage"] for student in final_response.json()["students"]}
        assert final_statuses[records[0]["student_id"]] == "absent"
        assert final_rates[records[0]["student_id"]] == 0
        assert await attendance.count_documents({
            "student_id": records[0]["student_id"],
            "course_id": course_id_str,
            "date": datetime.combine(date.fromisoformat(selected_date), datetime.min.time(), tzinfo=timezone.utc),
        }) == 1

        history_response = await client.get(
            f"/api/v1/attendance/course/{course_id_str}",
            params={"date": selected_date},
            headers=headers,
        )
        assert history_response.status_code == 200
        assert len(history_response.json()) == 3

        export_response = await client.get(
            "/api/v1/reports/attendance/excel",
            params={"course_id": course_id_str, "date": selected_date},
            headers=headers,
        )
        assert export_response.status_code == 200
        assert export_response.content.startswith(b"PK")
        from io import BytesIO
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(export_response.content), read_only=True)
        exported_rows = list(workbook.active.iter_rows(values_only=True))
        assert exported_rows[0] == ("Enrollment Number", "Student Name", "Date", "Status", "Course")
        assert any(row[1] == f"Attendance Test Student 1 {suffix}" and row[3] == "absent" for row in exported_rows[1:])

        unassigned_user_id = ObjectId()
        await users.insert_one({
            "_id": unassigned_user_id,
            "email": f"unassigned-{suffix}@test.sms.edu",
            "full_name": f"Unassigned Teacher {suffix}",
            "role": "teacher",
            "is_active": True,
        })
        unassigned_token = create_access_token({"sub": str(unassigned_user_id), "role": "teacher"})
        unassigned_courses = await client.get("/api/v1/attendance/courses", headers={"Authorization": f"Bearer {unassigned_token}"})
        assert unassigned_courses.status_code == 200
        assert any(course["code"] == "BCS301" for course in unassigned_courses.json())
        saved_by_unassigned_teacher = await client.post(
            "/api/v1/attendance",
            json={"course_id": course_id_str, "date": f"{selected_date}T00:00:00Z", "records": records},
            headers={"Authorization": f"Bearer {unassigned_token}"},
        )
        assert saved_by_unassigned_teacher.status_code == 200
        await users.delete_one({"_id": unassigned_user_id})
    finally:
        await attendance.delete_many({"course_id": course_id_str})
        await students.delete_many({"_id": {"$in": test_student_ids}})
        await users.delete_many({"_id": {"$in": [teacher_user_id, *test_student_user_ids]}})
        await teachers.delete_one({"_id": teacher_id})
        await courses.delete_one({"_id": course_id})


@pytest.mark.asyncio
async def test_cse_semester_three_student_discovers_curriculum_and_empty_attendance(client: AsyncClient):
    department = await get_collection("departments").find_one({"code": "CSE"})
    assert department
    suffix = uuid4().hex[:10]
    user_id = ObjectId()
    student_id = ObjectId()
    await get_collection("users").insert_one({
        "_id": user_id,
        "email": f"student-attendance-{suffix}@test.sms.edu",
        "full_name": f"CSE Semester Three Student {suffix}",
        "role": "student",
        "is_active": True,
    })
    await get_collection("students").insert_one({
        "_id": student_id,
        "user_id": str(user_id),
        "enrollment_number": f"CSE3{suffix.upper()}",
        "department_id": str(department["_id"]),
        "branch": "CSE",
        "semester": 3,
        "batch_year": 2025,
        "is_active": True,
    })
    headers = {"Authorization": f"Bearer {create_access_token({'sub': str(user_id), 'role': 'student'})}"}

    try:
        response = await client.get("/api/v1/attendance/student/courses", headers=headers)
        assert response.status_code == 200, response.text
        assert [course["code"] for course in response.json()] == [
            "BAS303", "BVE301", "BCS301", "BCS302", "BCS303", "BCC302",
            "BCS351", "BCS352", "BCS353", "BCC351",
        ]

        selected_course = next(course for course in response.json() if course["code"] == "BCS301")
        summary = await client.get(
            f"/api/v1/attendance/student/courses/{selected_course['id']}/summary",
            headers=headers,
        )
        assert summary.status_code == 200, summary.text
        assert summary.json() == {
            "course_id": selected_course["id"],
            "course_code": "BCS301",
            "course_name": "Data Structure",
            "present": 0,
            "absent": 0,
            "late": 0,
            "total_classes": 0,
            "attendance_percentage": 0.0,
            "has_records": False,
        }

        unrelated = await get_collection("courses").find_one({"code": "BAS101", "is_active": True})
        if unrelated:
            forbidden = await client.get(
                f"/api/v1/attendance/student/courses/{str(unrelated['_id'])}/summary",
                headers=headers,
            )
            assert forbidden.status_code == 403
    finally:
        await get_collection("students").delete_one({"_id": student_id})
        await get_collection("users").delete_one({"_id": user_id})
